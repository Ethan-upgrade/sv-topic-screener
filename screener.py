#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
多模态短视频内容筛选框架（clean-room 通用实现）
==================================================
用公开开源库实现“标题 + 口播语音转写 + 画面文字OCR”三路证据判定，
判断短视频是否属于某个主题（示例：汽车“智驾/辅助驾驶”内容）。
本工具不抓取任何平台、不内置任何第三方专有代码；素材由使用者自行合法提供。

输入：Excel（含 标题 / 链接 / 平台 列，列名支持常见别名）
可选：--media-dir 指向放有 <内容ID>.mp4 / <行号>.mp4 的本地视频目录；
      有视频则做语音转写 + 画面OCR，否则仅按文字判定。
输出：逐条结果 Excel。

用法：
  python screener.py --init-example                 # 生成示例 Excel
  python screener.py --input 数据.xlsx              # 仅文字判定
  python screener.py --input 数据.xlsx --media-dir 视频目录   # 文字+语音+画面
  python screener.py --input 数据.xlsx --rules rules.example.json
"""
from __future__ import annotations
import argparse, json, os, re, sys, time
from pathlib import Path

import openpyxl

TITLE_ALIASES = ["标题", "视频标题", "短视频标题", "作品标题", "笔记标题", "title"]
URL_ALIASES = ["链接", "视频链接", "作品链接", "笔记链接", "url", "url链接"]
PLATFORM_ALIASES = ["平台", "发布平台", "媒体平台", "platform"]

# 默认规则（示例：汽车“智驾/辅助驾驶”主题；可在 rules.json 覆盖）
DEFAULT_RULES = {
    "topic": "智驾",
    "smart_terms": [
        "智驾", "智能驾驶", "辅助驾驶", "自动驾驶", "高阶智驾", "领航辅助", "领航",
        "自动泊车", "记忆泊车", "代客泊车", "泊车", "激光雷达", "世界模型",
        "NOA", "NGP", "NCA", "ACC", "AEB", "APA", "AVP", "LCC", "LKA", "TJA",
    ],
    "brand_terms": ["Momenta"],
    "exclude_terms": ["纯话题标签", "无关样例词"],
    "count_brand_as_topic": False,
}


def _col(headers, aliases):
    low = [str(h).strip().lower() if h is not None else "" for h in headers]
    for i, h in enumerate(low):
        if h in [a.lower() for a in aliases]:
            return i
    return None


def read_rows(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    hdr = rows[0]
    ti = _col(hdr, TITLE_ALIASES)
    ui = _col(hdr, URL_ALIASES)
    pi = _col(hdr, PLATFORM_ALIASES)
    out = []
    for n, r in enumerate(rows[1:], start=2):
        title = str(r[ti]) if ti is not None and ti < len(r) and r[ti] is not None else ""
        url = str(r[ui]) if ui is not None and ui < len(r) and r[ui] is not None else ""
        plat = str(r[pi]) if pi is not None and pi < len(r) and r[pi] is not None else ""
        out.append({"row": n, "title": title, "url": url, "platform": plat})
    return hdr, out


def find_media(media_dir: Path, url: str, row: int):
    if not media_dir or not media_dir.exists():
        return None
    item_id = ""
    m = re.search(r"(?:video|note|explore|sph)/([A-Za-z0-9_-]{6,})", url or "")
    if m:
        item_id = m.group(1)
    for cand in ([item_id + ".mp4"] if item_id else []) + [f"{row}.mp4"]:
        p = media_dir / cand
        if p.exists() and p.stat().st_size > 0:
            return p
    return None


_whisper = None


def transcribe(path: Path):
    """口播语音转写（faster-whisper 自带 PyAV 解码，无需系统 ffmpeg）。"""
    global _whisper
    try:
        from faster_whisper import WhisperModel
        if _whisper is None:
            _whisper = WhisperModel(os.environ.get("WHISPER_MODEL", "small"), device="cpu", compute_type="int8")
        segs, info = _whisper.transcribe(str(path), vad_filter=True, beam_size=1)
        return "".join(s.text for s in segs).strip()
    except Exception as e:
        return ""


_ocr = None


def ocr_frames(video: Path, n_frames: int = 6):
    """画面文字 OCR：用 OpenCV 均匀抽 n 帧，RapidOCR 识别，合并文本。"""
    global _ocr
    texts = []
    try:
        import cv2
        cap = cv2.VideoCapture(str(video))
        total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT)) or 0
        if total > 0:
            from rapidocr import RapidOCR
            if _ocr is None:
                _ocr = RapidOCR()
            for k in range(n_frames):
                idx = int(total * (k + 0.5) / n_frames)
                cap.set(cv2.CAP_PROP_POS_FRAMES, idx)
                ok, frame = cap.read()
                if not ok or frame is None:
                    continue
                res = _ocr(frame)
                t = ""
                try:
                    if hasattr(res, "txts"):
                        t = "\n".join([x for x in (res.txts or [])])
                    elif isinstance(res, tuple) and len(res) >= 2:
                        t = "\n".join(res[1] or [])
                except Exception:
                    t = ""
                if t:
                    texts.append(t)
        cap.release()
    except Exception:
        pass
    return "\n".join(texts).strip()


def hits(text, terms):
    t = (text or "").lower()
    return [k for k in terms if k.lower() in t]


def decide(rec, rules, media_dir, do_media=True):
    text = rec["title"]
    terms = list(rules.get("smart_terms", []))
    if rules.get("count_brand_as_topic"):
        terms += list(rules.get("brand_terms", []))
    h = hits(text, terms)
    if h:
        return "相关", "文本(标题/话题)", h, "", "", "done_text"
    if not do_media:
        return "待人工复核", "文本未命中", [], "", "", "boundary_no_media"
    vid = find_media(media_dir, rec["url"], rec["row"])
    if not vid:
        return "待人工复核", "无本地视频", [], "", "", "boundary_no_media"
    tr = transcribe(vid)
    oc = ocr_frames(vid)
    all_terms = terms + list(rules.get("brand_terms", []))
    mh = list(dict.fromkeys(hits(tr, all_terms) + hits(oc, all_terms)))
    if mh:
        return "相关", "媒体(口播/画面)", mh, tr[:200], oc[:200], "done_media"
    return "不相关", "媒体均未命中", [], tr[:200], oc[:200], "done_media"


def init_example(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["平台", "视频链接", "短视频标题"])
    ws.append(["抖音", "https://example.com/video/1000000000000000001", "MG07 搭载 Momenta 高阶智驾 激光雷达 自动泊车"])
    ws.append(["抖音", "https://example.com/video/1000000000000000002", "今天天气不错，随便拍拍生活碎片"])
    wb.save(path)
    print("已生成示例：", path)


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", default="")
    ap.add_argument("--out", default="")
    ap.add_argument("--media-dir", default="")
    ap.add_argument("--rules", default="")
    ap.add_argument("--no-media", action="store_true", help="仅按文字判定（不做语音/画面）")
    ap.add_argument("--init-example", action="store_true")
    args = ap.parse_args(argv)

    if args.init_example:
        init_example(Path(args.input or "example_input.xlsx"))
        return 0
    if not args.input:
        ap.error("需要 --input 指定 Excel（或先用 --init-example 生成示例）")

    inp = Path(args.input)
    out = Path(args.out) if args.out else inp.with_name(inp.stem + "_筛选结果.xlsx")
    rules = dict(DEFAULT_RULES)
    if args.rules and Path(args.rules).exists():
        rules.update(json.loads(Path(args.rules).read_text(encoding="utf-8")))
    media_dir = Path(args.media_dir) if args.media_dir else None

    hdr, rows = read_rows(inp)
    print("记录数:", len(rows), "| 主题:", rules.get("topic"))
    results = []
    t0 = time.time()
    for i, rec in enumerate(rows, 1):
        dec, src, hs, tr, oc, status = decide(rec, rules, media_dir, do_media=not args.no_media)
        results.append([i, rec["platform"], rec["title"], rec["url"], dec, src, "、".join(hs[:8]), tr, oc, status])
        if i % 50 == 0:
            print("  进度 %d/%d" % (i, len(rows)))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "筛选结果"
    cols = ["序号", "平台", "标题", "链接", "判定", "命中来源", "命中词", "口播片段", "画面片段", "状态"]
    ws.append(cols)
    for r in results:
        ws.append(r)
    from collections import Counter
    dist = Counter(r[4] for r in results)
    ws2 = wb.create_sheet("统计")
    ws2.append(["判定", "条数"])
    for k, v in dist.items():
        ws2.append([k, v])
    ws2.append(["合计", len(results)])
    ws2.append(["耗时秒", round(time.time() - t0, 1)])
    wb.save(out)
    print("完成:", len(results), "条，用时 %.1fs" % (time.time() - t0))
    print("判定分布:", dict(dist))
    print("结果文件:", out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
